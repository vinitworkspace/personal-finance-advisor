"""
Script to compile all markdown documentation files into a professional,
publication-quality Microsoft Word (.docx) project report.
"""

import os
import re
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = Path(__file__).resolve().parent
DOCS_DIR = BASE_DIR / "docs"
OUTPUT_DOCX = DOCS_DIR / "Personal_Finance_Advisor_Bot_Documentation.docx"

# Color Palette
COLOR_PRIMARY = RGBColor(15, 23, 42)      # Deep Slate #0F172A
COLOR_SECONDARY = RGBColor(37, 99, 235)   # Tech Blue #2563EB
COLOR_ACCENT = RGBColor(13, 148, 136)     # Emerald/Teal #0D9488
COLOR_BODY = RGBColor(51, 65, 85)         # Charcoal Slate #334155
COLOR_MUTED = RGBColor(100, 116, 139)     # Cool Gray #64748B
HEX_HEADER_BG = "1E293B"                  # Dark table header #1E293B
HEX_ALT_ROW = "F8FAFC"                    # Light tint row #F8FAFC
HEX_CODE_BG = "F1F5F9"                    # Light code block #F1F5F9

def set_cell_background(cell, hex_color: str):
    """Set background color of a table cell."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner cell padding."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_title_cover_page(doc):
    """Generate a clean, high-impact project title page."""
    # Top spacing
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(60)

    # Project Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("PERSONAL FINANCE ADVISOR BOT")
    r_title.bold = True
    r_title.font.name = "Segoe UI"
    r_title.font.size = Pt(28)
    r_title.font.color.rgb = COLOR_PRIMARY
    p_title.paragraph_format.space_after = Pt(10)

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Full-Stack Intelligent Financial Planning, Budgeting & Advisory System")
    r_sub.font.name = "Segoe UI"
    r_sub.font.size = Pt(14)
    r_sub.font.color.rgb = COLOR_SECONDARY
    p_sub.paragraph_format.space_after = Pt(20)

    # Badge/Tagline
    p_badge = doc.add_paragraph()
    p_badge.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_badge = p_badge.add_run("Comprehensive Software Architecture, Security, Math & Viva Voce Manual")
    r_badge.font.name = "Segoe UI"
    r_badge.font.size = Pt(11)
    r_badge.italic = True
    r_badge.font.color.rgb = COLOR_MUTED
    p_badge.paragraph_format.space_after = Pt(80)

    # Project Metadata Box (2-column key-value table)
    meta_table = doc.add_table(rows=6, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Technology Stack", "Python 3.11, Flask 3.0, SQLite 3, SQLAlchemy 2.0, Chart.js 4.4"),
        ("Artificial Intelligence", "Google Gemini 1.5 Flash API + Offline Heuristic Rule Engine"),
        ("Architecture Pattern", "4-Tier MVC / Blueprint / Service Layer Pattern"),
        ("Target Evaluation", "B.Sc. Computer Science / Cybersecurity Academic Project & Viva"),
        ("Date of Publication", "March 2026"),
        ("Version & Status", "Version 1.0.0 (Production-Grade, 42/42 Tests Passing)"),
    ]

    for idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[idx]
        
        c0 = row.cells[0]
        c0.width = Inches(2.2)
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(label)
        r0.bold = True
        r0.font.name = "Segoe UI"
        r0.font.size = Pt(10)
        r0.font.color.rgb = COLOR_PRIMARY
        p0.paragraph_format.space_after = Pt(2)
        set_cell_background(c0, "F8FAFC")
        set_cell_margins(c0, top=80, bottom=80, left=120, right=120)

        c1 = row.cells[1]
        c1.width = Inches(4.3)
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.name = "Segoe UI"
        r1.font.size = Pt(10)
        r1.font.color.rgb = COLOR_BODY
        p1.paragraph_format.space_after = Pt(2)
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c1, top=80, bottom=80, left=120, right=120)

    # Page Break after cover
    doc.add_page_break()

def add_header_footer(doc):
    """Add uniform running headers and footers."""
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

        # Header
        header = section.header
        p_head = header.paragraphs[0]
        p_head.text = "Personal Finance Advisor Bot  |  Technical Documentation"
        p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_head.runs[0].font.name = "Segoe UI"
        p_head.runs[0].font.size = Pt(8.5)
        p_head.runs[0].font.color.rgb = COLOR_MUTED

        # Footer
        footer = section.footer
        p_foot = footer.paragraphs[0]
        p_foot.text = "Personal Finance Advisor Bot — Production Quality Academic Project"
        p_foot.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_foot.runs[0].font.name = "Segoe UI"
        p_foot.runs[0].font.size = Pt(8.5)
        p_foot.runs[0].font.color.rgb = COLOR_MUTED

def format_inline_runs(paragraph, text: str):
    """Parse inline bold (`**word**`), code (`word`), and italic (`*word*`)."""
    # Regex pattern to identify markdown inline tokens
    tokens = re.split(r'(\*\*.*?\*\*|`.*?`|\*.*?\*)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) >= 4:
            r = paragraph.add_run(token[2:-2])
            r.bold = True
            r.font.name = "Segoe UI"
            r.font.size = Pt(10)
            r.font.color.rgb = COLOR_PRIMARY
        elif token.startswith("`") and token.endswith("`") and len(token) >= 2:
            r = paragraph.add_run(token[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(194, 65, 12)  # Dark amber/orange code color
        elif token.startswith("*") and token.endswith("*") and len(token) >= 2:
            r = paragraph.add_run(token[1:-1])
            r.italic = True
            r.font.name = "Segoe UI"
            r.font.size = Pt(10)
            r.font.color.rgb = COLOR_BODY
        else:
            r = paragraph.add_run(token)
            r.font.name = "Segoe UI"
            r.font.size = Pt(10)
            r.font.color.rgb = COLOR_BODY

def add_callout_box(doc, text: str):
    """Add a stylized callout / alert box."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    cell.width = Inches(6.5)
    set_cell_background(cell, "F0FDF4")  # Soft emerald tint
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    format_inline_runs(p, text)

def add_code_block(doc, code_lines: list):
    """Add shaded preformatted code block."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    cell.width = Inches(6.5)
    set_cell_background(cell, HEX_CODE_BG)
    set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    
    r = p.add_run("\n".join(code_lines))
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(30, 41, 59)

def add_markdown_table(doc, table_lines: list):
    """Render a markdown table into a beautifully formatted docx table."""
    rows_data = []
    for line in table_lines:
        line = line.strip()
        if not line or not line.startswith("|"):
            continue
        # Split by pipe and remove empty boundary items
        cells = [c.strip() for c in line.split("|")[1:-1]]
        # Skip markdown separator row like |:---|:---|
        if cells and all(set(c).issubset({"-", ":", " "}) for c in cells):
            continue
        rows_data.append(cells)

    if not rows_data:
        return

    num_cols = max(len(r) for r in rows_data)
    tbl = doc.add_table(rows=len(rows_data), cols=num_cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    col_width = Inches(6.5 / max(1, num_cols))

    for r_idx, row_cells in enumerate(rows_data):
        row = tbl.rows[r_idx]
        is_header = (r_idx == 0)
        
        for c_idx in range(num_cols):
            cell = row.cells[c_idx]
            cell.width = col_width
            cell_val = row_cells[c_idx] if c_idx < len(row_cells) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)

            if is_header:
                set_cell_background(cell, HEX_HEADER_BG)
                r = p.add_run(cell_val)
                r.bold = True
                r.font.name = "Segoe UI"
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(255, 255, 255)
            else:
                bg = HEX_ALT_ROW if r_idx % 2 == 1 else "FFFFFF"
                set_cell_background(cell, bg)
                format_inline_runs(p, cell_val)

            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

def parse_markdown_to_docx(doc, filepath: Path, chapter_title: str):
    """Parse a single markdown file and append formatted elements to docx."""
    # Section Heading
    p_chap = doc.add_paragraph()
    p_chap.paragraph_format.space_before = Pt(18)
    p_chap.paragraph_format.space_after = Pt(8)
    p_chap.paragraph_format.keep_with_next = True
    r_chap = p_chap.add_run(chapter_title)
    r_chap.bold = True
    r_chap.font.name = "Segoe UI"
    r_chap.font.size = Pt(18)
    r_chap.font.color.rgb = COLOR_PRIMARY

    with open(filepath, "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    for raw_line in lines:
        line = raw_line.rstrip()
        stripped = line.strip()

        # Handle Code Blocks (```)
        if stripped.startswith("```"):
            if in_code_block:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code_block = False
            else:
                if in_table:
                    add_markdown_table(doc, table_lines)
                    table_lines = []
                    in_table = False
                in_code_block = True
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        # Handle Tables (| col1 | col2 |)
        if stripped.startswith("|") and stripped.endswith("|"):
            in_table = True
            table_lines.append(stripped)
            continue
        elif in_table:
            add_markdown_table(doc, table_lines)
            table_lines = []
            in_table = False

        # Blank line
        if not stripped:
            continue

        # Skip primary file title since chapter_title is added
        if stripped.startswith("# "):
            continue

        # Heading 2 (## ...)
        if stripped.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(stripped[3:])
            r.bold = True
            r.font.name = "Segoe UI"
            r.font.size = Pt(13.5)
            r.font.color.rgb = COLOR_SECONDARY
            continue

        # Heading 3 (### ...)
        if stripped.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(stripped[4:])
            r.bold = True
            r.font.name = "Segoe UI"
            r.font.size = Pt(11)
            r.font.color.rgb = COLOR_PRIMARY
            continue

        # Blockquote (> ...)
        if stripped.startswith("> "):
            add_callout_box(doc, stripped[2:])
            continue

        # Bullet List (- ... or * ...)
        if stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Inches(0.25)
            format_inline_runs(p, stripped[2:])
            continue

        # Numbered List (1. ... or 2. ...)
        match_num = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if match_num:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Inches(0.25)
            format_inline_runs(p, match_num.group(2))
            continue

        # Standard Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        format_inline_runs(p, stripped)

    # Flush any remaining buffer
    if in_code_block and code_lines:
        add_code_block(doc, code_lines)
    if in_table and table_lines:
        add_markdown_table(doc, table_lines)

    # Page break after chapter
    doc.add_page_break()

def generate_master_docx():
    """Compile all documentation chapters into a unified Word document."""
    print("=" * 65)
    print(" Generating Microsoft Word Project Documentation (.docx)...")
    print("=" * 65)

    doc = docx.Document()
    add_header_footer(doc)
    add_title_cover_page(doc)

    chapters = [
        ("01_ARCHITECTURE_AND_DESIGN.md", "Chapter 1: System Architecture & Software Design"),
        ("02_DATABASE_AND_MODELS.md", "Chapter 2: Database Architecture & Data Models"),
        ("03_SECURITY_AND_AUTHENTICATION.md", "Chapter 3: Cybersecurity, Authentication & Tenant Isolation"),
        ("04_AI_ENGINE_AND_FALLBACK.md", "Chapter 4: Artificial Intelligence, Prompt Engineering & Fallback"),
        ("05_FINANCIAL_MATHEMATICS_AND_FORMULAS.md", "Chapter 5: Financial Mathematics & Algorithmic Calculations"),
        ("06_API_REFERENCE.md", "Chapter 6: Complete REST API Reference & Payload Specifications"),
        ("07_HOW_TO_RUN_AND_BROWSER_GUIDE.md", "Chapter 7: How to Run the Application & Complete Browser Tour"),
        ("08_VIVA_QUESTIONS_AND_ANSWERS.md", "Chapter 8: Viva Voce Examination Questions & Model Answers"),
        ("09_VERCEL_DEPLOYMENT_GUIDE.md", "Chapter 9: Vercel Serverless Cloud Deployment Guide"),
    ]

    for filename, title in chapters:
        filepath = DOCS_DIR / filename
        if filepath.exists():
            print(f" [+] Compiling: {title} ...")
            parse_markdown_to_docx(doc, filepath, title)
        else:
            print(f" [!] Missing file: {filepath}")

    doc.save(str(OUTPUT_DOCX))
    print(f"\n[SUCCESS] Document saved to: {OUTPUT_DOCX}")
    print(f"File size: {os.path.getsize(OUTPUT_DOCX) / 1024:.1f} KB")

if __name__ == "__main__":
    generate_master_docx()
